"""Lectura de documentos Word (.docx).

Basado en el extractor del Anonimizador Judicial de IALAB — Facultad de
Derecho, UBA (Apache 2.0). Modificado: ahora se recorren siempre las tablas,
los encabezados y los pies de página, porque en los documentos judiciales
chilenos es frecuente que el RIT, el tribunal y la individualización de las
partes aparezcan justamente ahí y no en el cuerpo.
"""
from __future__ import annotations

import io

from docx import Document
from docx.table import Table


def _texto_de_tabla(tabla: Table) -> list[str]:
    filas: list[str] = []
    for fila in tabla.rows:
        celdas = [celda.text.strip() for celda in fila.cells if celda.text.strip()]
        # Se eliminan las repeticiones que produce python-docx cuando una celda
        # está combinada horizontalmente.
        unicas: list[str] = []
        for celda in celdas:
            if not unicas or unicas[-1] != celda:
                unicas.append(celda)
        if unicas:
            filas.append(" | ".join(unicas))
    return filas


def extract_docx(data: bytes) -> str:
    """Devuelve el texto del documento Word, incluidas tablas y encabezados."""
    documento = Document(io.BytesIO(data))

    partes: list[str] = []

    for seccion in documento.sections:
        for contenedor in (seccion.header, seccion.footer):
            for parrafo in contenedor.paragraphs:
                if parrafo.text.strip():
                    partes.append(parrafo.text.strip())

    for parrafo in documento.paragraphs:
        if parrafo.text.strip():
            partes.append(parrafo.text.strip())

    for tabla in documento.tables:
        partes.extend(_texto_de_tabla(tabla))

    texto = "\n".join(dict.fromkeys(partes)).strip()
    if not texto:
        raise ValueError(
            "No se pudo extraer texto del documento Word. Es posible que su "
            "contenido sean imágenes; en ese caso, esta aplicación no puede "
            "procesarlo porque no incluye reconocimiento óptico de caracteres."
        )
    return texto
