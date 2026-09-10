"""Formato del documento exportado.

Basado en el módulo de formato del Anonimizador Judicial de IALAB — Facultad
de Derecho, UBA (Apache 2.0). Modificado: la numeración de páginas pasó de
romana a arábiga, que es la de uso corriente en los escritos chilenos, y los
márgenes por defecto se ajustaron a esa práctica.
"""
from __future__ import annotations

from dataclasses import dataclass

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


@dataclass(frozen=True)
class FormatoDocumento:
    font_name: str = "Times New Roman"
    font_size_pt: int = 12
    line_spacing: float = 1.5
    space_after_pt: int = 6
    margin_cm: float = 3.0
    margin_top_bottom_cm: float = 2.5
    page_number_format: str = "decimal"
    alignment: str = "justify"


FORMATO_POR_DEFECTO = FormatoDocumento()

_ALINEACIONES = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def _aplicar_fuente(fuente, formato: FormatoDocumento) -> None:
    fuente.name = formato.font_name
    fuente.size = Pt(formato.font_size_pt)


def _forzar_fuente_en_fragmento(fragmento, formato: FormatoDocumento) -> None:
    """Fija la tipografía en el XML para que Word no la sustituya."""
    _aplicar_fuente(fragmento.font, formato)
    propiedades = fragmento._element.get_or_add_rPr()
    fuentes = propiedades.rFonts
    if fuentes is None:
        fuentes = OxmlElement("w:rFonts")
        propiedades.insert(0, fuentes)
    for atributo in ("ascii", "hAnsi", "cs", "eastAsia"):
        fuentes.set(qn(f"w:{atributo}"), formato.font_name)


def _aplicar_formato_parrafo(parrafo: Paragraph, formato: FormatoDocumento) -> None:
    propiedades = parrafo.paragraph_format
    propiedades.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    propiedades.line_spacing = formato.line_spacing
    propiedades.space_after = Pt(formato.space_after_pt)
    propiedades.alignment = _ALINEACIONES.get(
        formato.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY
    )
    for fragmento in parrafo.runs:
        _forzar_fuente_en_fragmento(fragmento, formato)


def _configurar_estilo_normal(documento: Document, formato: FormatoDocumento) -> None:
    estilo = documento.styles["Normal"]
    _aplicar_fuente(estilo.font, formato)
    propiedades = estilo.paragraph_format
    propiedades.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    propiedades.line_spacing = formato.line_spacing
    propiedades.space_after = Pt(formato.space_after_pt)
    propiedades.alignment = _ALINEACIONES.get(
        formato.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY
    )


def _configurar_pagina(documento: Document, formato: FormatoDocumento) -> None:
    seccion = documento.sections[0]
    margen = Cm(formato.margin_cm)
    seccion.left_margin = margen
    seccion.right_margin = margen
    seccion.top_margin = Cm(formato.margin_top_bottom_cm)
    seccion.bottom_margin = Cm(formato.margin_top_bottom_cm)
    seccion.gutter = Cm(0)


def _agregar_campo_numero_pagina(parrafo: Paragraph) -> None:
    fragmento = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instruccion = OxmlElement("w:instrText")
    instruccion.set(qn("xml:space"), "preserve")
    instruccion.text = "PAGE"
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    elemento = fragmento._r
    elemento.append(inicio)
    elemento.append(instruccion)
    elemento.append(separador)
    elemento.append(fin)


def _configurar_numeracion(documento: Document, formato: FormatoDocumento) -> None:
    seccion = documento.sections[0]
    propiedades = seccion._sectPr
    numeracion = propiedades.find(qn("w:pgNumType"))
    if numeracion is None:
        numeracion = OxmlElement("w:pgNumType")
        propiedades.append(numeracion)
    numeracion.set(qn("w:fmt"), formato.page_number_format)
    numeracion.set(qn("w:start"), "1")

    pie = seccion.footer
    pie.is_linked_to_previous = False
    parrafo = pie.paragraphs[0] if pie.paragraphs else pie.add_paragraph()
    parrafo.clear()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _agregar_campo_numero_pagina(parrafo)


def aplicar_formato(documento: Document, formato: FormatoDocumento | None = None) -> None:
    """Aplica márgenes, tipografía, interlineado y numeración al documento."""
    formato = formato or FORMATO_POR_DEFECTO
    _configurar_pagina(documento, formato)
    _configurar_estilo_normal(documento, formato)
    _configurar_numeracion(documento, formato)


def agregar_parrafo(
    documento: Document, texto: str, formato: FormatoDocumento | None = None
) -> Paragraph:
    formato = formato or FORMATO_POR_DEFECTO
    parrafo = documento.add_paragraph(texto)
    _aplicar_formato_parrafo(parrafo, formato)
    return parrafo
