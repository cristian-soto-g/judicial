"""Pruebas de la lectura de documentos."""
from __future__ import annotations

import io

import pytest
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.extraction import extraer_documento
from app.extraction.pdf_cleanup import clean_pdf_text, merge_lines_to_paragraphs


def _docx_de_prueba(con_tabla: bool = False) -> bytes:
    documento = Document()
    documento.add_paragraph("Sentencia definitiva en causa RIT 145-2024.")
    if con_tabla:
        tabla = documento.add_table(rows=2, cols=2)
        tabla.cell(0, 0).text = "Imputado"
        tabla.cell(0, 1).text = "Juan Pérez Muñoz"
        tabla.cell(1, 0).text = "RUT"
        tabla.cell(1, 1).text = "12.345.678-5"
    buffer = io.BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


def _pdf_con_texto(texto: str) -> bytes:
    buffer = io.BytesIO()
    lienzo = canvas.Canvas(buffer, pagesize=A4)
    lienzo.drawString(70, 780, texto)
    lienzo.save()
    return buffer.getvalue()


def _pdf_sin_texto() -> bytes:
    """PDF que solo contiene un dibujo: simula una página escaneada."""
    buffer = io.BytesIO()
    lienzo = canvas.Canvas(buffer, pagesize=A4)
    lienzo.rect(70, 700, 400, 80, fill=1)
    lienzo.save()
    return buffer.getvalue()


class TestTextoPlano:
    @pytest.mark.parametrize("codificacion", ["utf-8", "cp1252", "latin-1"])
    def test_lee_las_codificaciones_habituales(self, codificacion):
        contenido = "Compareció don Juan Pérez Muñoz.".encode(codificacion)
        resultado = extraer_documento("nota.txt", contenido)
        assert "Juan Pérez Muñoz" in resultado.texto

    def test_normaliza_los_saltos_de_linea(self):
        resultado = extraer_documento("nota.txt", b"Primera\r\nSegunda\r\n")
        assert resultado.texto == "Primera\nSegunda"

    def test_rechaza_archivo_vacio(self):
        with pytest.raises(ValueError):
            extraer_documento("nota.txt", b"   ")


class TestWord:
    def test_lee_los_parrafos(self):
        resultado = extraer_documento("causa.docx", _docx_de_prueba())
        assert "RIT 145-2024" in resultado.texto

    def test_lee_las_tablas(self):
        # En los documentos judiciales chilenos la individualización de las
        # partes viaja con frecuencia dentro de una tabla.
        resultado = extraer_documento("causa.docx", _docx_de_prueba(con_tabla=True))
        assert "Juan Pérez Muñoz" in resultado.texto
        assert "12.345.678-5" in resultado.texto


class TestPdf:
    def test_lee_el_texto_seleccionable(self):
        resultado = extraer_documento("causa.pdf", _pdf_con_texto("RUT 12.345.678-5"))
        assert "12.345.678-5" in resultado.texto

    def test_rechaza_el_pdf_escaneado_con_un_mensaje_claro(self):
        # No basta con fallar: el mensaje debe explicar por qué, porque
        # procesar un escaneo en silencio produciría un documento en apariencia
        # anonimizado y en realidad intacto.
        with pytest.raises(ValueError) as error:
            extraer_documento("escaneo.pdf", _pdf_sin_texto())
        mensaje = str(error.value).lower()
        assert "escaneo" in mensaje or "seleccionable" in mensaje


class TestFormatosNoAdmitidos:
    def test_explica_como_convertir_el_formato_antiguo_de_word(self):
        with pytest.raises(ValueError) as error:
            extraer_documento("causa.doc", b"contenido")
        assert ".docx" in str(error.value)

    def test_rechaza_otros_formatos(self):
        with pytest.raises(ValueError):
            extraer_documento("imagen.png", b"contenido")


class TestLimpiezaDePdf:
    def test_quita_el_sello_de_firma_electronica(self):
        texto = (
            "Resuelvo: ha lugar.\n"
            "Este documento tiene firma electrónica avanzada\n"
            "Código de verificación: ABC123\n"
            "Notifíquese."
        )
        limpio = clean_pdf_text(texto)
        assert "firma electrónica" not in limpio.lower()
        assert "Notifíquese" in limpio

    def test_une_las_lineas_de_un_mismo_parrafo(self):
        texto = "El tribunal resolvió que la\nsolicitud es procedente."
        assert merge_lines_to_paragraphs(texto).count("\n") == 0
